"""Convert a Markdown document with LaTeX equations to Word, with the
equations as native (editable) Word equations.

    python3 tools/md_math_to_docx.py APPENDIX.md APPENDIX.docx

Everything `md_to_docx.py` handles, plus:

- `$...$` inline math and `$$...$$` display math (a display equation is a
  line of its own), converted LaTeX -> MathML (latex2mathml) -> OMML with the
  stylesheet Word ships (`mathml2omml.xsl`);
- `*italic*`;
- the first `# ` heading becomes the document title.

`md_to_docx.py` itself is untouched, so MEMO.docx builds exactly as before.
If the stylesheet or latex2mathml is missing, equations fall back to their
LaTeX source in italics, and the script says so.
"""

from __future__ import annotations

import os
import re
import sys

from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import md_to_docx as base  # noqa: E402

XSL_CANDIDATES = [
    "/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl",
    "/Applications/Microsoft Word.app/Contents/Resources/MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
]

_transform = None
_warned = False


_SPACES = {"\\qquad": "\u2003\u2003\u2003", "\\quad": "\u2003", "\\;": "\u2005",
           "\\,": "\u2009", "\\ ": "\u2005"}
_VERBATIM = {"\\text", "\\mathrm", "\\mathit", "\\operatorname", "\\begin", "\\end"}


def _prep(latex):
    """Make LaTeX survive the MathML->OMML round trip the way a reader expects:

    - spacing commands become literal space characters inside \\text (the
      stylesheet drops <mspace>), so equation numbers keep their distance;
    - a run of two or more letters is one name (\\mathit{QM}), not a product
      of single letters, so `PWM_c` subscripts the name, not its last letter;
    - whitespace between two names in the source becomes a thin space, as a
      typesetter would leave between juxtaposed variables.
    """
    out, i, n = [], 0, len(latex)
    prev_atom = False                     # last emitted thing was a name/number/group end
    while i < n:
        ch = latex[i]
        if ch == "\\":
            m = re.match(r"\\([A-Za-z]+|.)", latex[i:])
            cmd = m.group(0)
            i += len(cmd)
            if cmd in _SPACES:
                out.append("\\text{" + _SPACES[cmd] + "}")
                prev_atom = False
                continue
            out.append(cmd)
            if cmd in _VERBATIM and i < n and latex[i] == "{":
                depth, j = 0, i
                while j < n:
                    depth += {"{": 1, "}": -1}.get(latex[j], 0)
                    j += 1
                    if depth == 0:
                        break
                out.append(latex[i:j])
                i = j
                prev_atom = cmd not in ("\\begin", "\\end")
            else:
                prev_atom = False
            continue
        if ch.isspace():
            j = i
            while j < n and latex[j].isspace():
                j += 1
            if prev_atom and j < n and latex[j].isalpha():
                out.append("\\text{\u2009}")
            i = j
            continue
        m = re.match(r"[A-Za-z]{2,}", latex[i:])
        if m:
            out.append("\\mathit{" + m.group(0) + "}")
            i += len(m.group(0))
            prev_atom = True
            continue
        out.append(ch)
        prev_atom = ch.isalnum() or ch in "})"
        i += 1
    return "".join(out)


def _omml(latex, display=False):
    """OMML element for a LaTeX string, or None if conversion is unavailable."""
    global _transform, _warned
    try:
        from latex2mathml.converter import convert as l2m
        from lxml import etree
    except ImportError:
        _transform = False
    if _transform is None:
        path = next((p for p in XSL_CANDIDATES if os.path.exists(p)), None)
        _transform = etree.XSLT(etree.parse(path)) if path else False
    if not _transform:
        if not _warned:
            print("!! no MathML->OMML stylesheet or latex2mathml: equations left as LaTeX text")
            _warned = True
        return None
    mml = l2m(_prep(latex), display="block" if display else "inline")
    return _transform(etree.fromstring(mml)).getroot()


def _add_math(par, latex, display=False):
    el = _omml(latex, display)
    if el is None:
        r = par.add_run(latex)
        r.italic = True
        return
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    if display and el.tag != f"{{{ns}}}oMathPara":
        from lxml import etree
        para = etree.SubElement(par._p, f"{{{ns}}}oMathPara")
        para.append(el)
    elif not display and el.tag == f"{{{ns}}}oMathPara":
        for child in list(el):              # inline: unwrap to the oMath
            par._p.append(child)
    else:
        par._p.append(el)


TOKEN = re.compile(r"(\[[^\]]+\]\([^)\s]+\)|\$[^$]+\$|\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def _add_link(par, text, url):
    """A clickable hyperlink run (blue, underlined)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F5C99")
    under = OxmlElement("w:u"); under.set(qn("w:val"), "single")
    rpr.append(color); rpr.append(under); run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); link.append(run)
    par._p.append(link)


def add_runs(par, text):
    for piece in TOKEN.split(text):
        if not piece:
            continue
        m = LINK.fullmatch(piece)
        if m:
            _add_link(par, m.group(1), m.group(2))
        elif piece.startswith("$") and piece.endswith("$") and len(piece) > 2:
            _add_math(par, piece[1:-1])
        elif piece.startswith("**") and piece.endswith("**"):
            _bold(par, piece[2:-2])
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = base.CODE_FONT
            r.font.size = Pt(9.5)
            r.font.color.rgb = base.MONO_COLOR
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def _bold(par, text):
    """Bold text that may itself hold *italic* or $math$."""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("$") and piece.endswith("$") and len(piece) > 2:
            _add_math(par, piece[1:-1])
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.bold = r.italic = True
        else:
            par.add_run(piece).bold = True


def emit_table(doc, rows):
    """md_to_docx's table, with column widths in proportion to their text
    (equal widths squeeze a long description column into a sliver)."""
    from docx.shared import Inches
    base_emit(doc, rows)
    table = doc.tables[-1]
    ncol = len(table.columns)
    usable = 6.5                                         # inches, Letter/A4 with default margins
    mins, weight = [], []
    for j in range(ncol):
        cells = [r[j] if j < len(r) else "" for r in rows]
        clean = [re.sub(r"[`*$\\{}]", "", c) for c in cells]
        avg = sum(len(c) for c in clean[1:]) / max(len(clean) - 1, 1)
        # Word can break after a hyphen, dash or slash, so those split words
        longest_word = max((len(w) for c in clean for w in re.split(r"[\s\-–/]+", c)), default=4)
        mins.append(0.065 * longest_word + 0.18)     # the longest word fits unbroken
        weight.append(max(avg, 1.0))
    spare = max(usable - sum(mins), 0.0)
    widths = [m + spare * w / sum(weight) for m, w in zip(mins, weight)]
    scale = min(1.0, usable / sum(widths))
    table.autofit = False
    for j, w in enumerate(widths):
        for cell in table.columns[j].cells:
            cell.width = Inches(w * scale)


base_emit = base.emit_table


def convert(src, dst, public=False):
    base.add_runs = add_runs              # tables and lists go through it too
    base.emit_table = emit_table
    from doc_variant import variant       # internal / public passages (tools/doc_variant.py)
    text = variant(open(src, encoding="utf-8").read(), public=public)
    doc = base.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = text.splitlines()
    i, para, table, code = 0, [], [], None
    titled = False

    def flush_para():
        if para:
            add_runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    def flush_table():
        if table:
            base.emit_table(doc, list(table))
            table.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if code is not None:
            if stripped.startswith("```"):
                base.emit_code(doc, code)
                code = None
            else:
                code.append(line)
            i += 1
            continue
        if stripped.startswith("```"):
            flush_para(); flush_table()
            code = []
        elif stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
            flush_para(); flush_table()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            _add_math(p, stripped[2:-2].strip(), display=True)
        elif base.is_table_row(stripped):
            flush_para()
            cells = base.split_row(stripped)
            if not base.is_divider(cells):
                table.append(cells)
        else:
            flush_table()
            if not stripped:
                flush_para()
            elif stripped.startswith("#"):
                flush_para()
                level = len(stripped) - len(stripped.lstrip("#"))
                title = re.sub(r"[`*]", "", stripped.lstrip("# ").strip())
                if level == 1 and not titled:
                    doc.add_heading(title, level=0)
                    titled = True
                else:
                    doc.add_heading(title, level=min(level, 4))
            elif re.fullmatch(r"-{3,}", stripped):
                flush_para()
            elif stripped.startswith(("- ", "* ")):
                flush_para()
                add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
            elif re.match(r"\d+\.\s", stripped):
                flush_para()
                add_runs(doc.add_paragraph(style="List Number"),
                         re.sub(r"^\d+\.\s+", "", stripped))
            else:
                para.append(stripped)
        i += 1

    flush_para(); flush_table()
    doc.save(dst)
    return dst


if __name__ == "__main__":
    # --public: the public version (tools/doc_variant.py), for the public repository
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = args[0] if args else "APPENDIX.md"
    dst = args[1] if len(args) > 1 else os.path.splitext(src)[0] + ".docx"
    print("wrote", convert(src, dst, public="--public" in sys.argv))
