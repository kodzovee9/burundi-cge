"""Convert the working memo from Markdown to Word.

    python3 tools/md_to_docx.py MEMO.md MEMO.docx

Handles the constructs the memo actually uses -- headings, paragraphs, pipe
tables, fenced code, bullet and numbered lists, and inline bold/code -- rather
than trying to be a general Markdown implementation. Re-run it whenever MEMO.md
changes so the two stay in step.
"""

from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

CODE_FONT = "Menlo"
MONO_COLOR = RGBColor(0x8A, 0x1C, 0x1C)


LINK = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def add_link(par, text, url):
    """A clickable hyperlink run (blue, underlined)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F5C99")
    under = OxmlElement("w:u"); under.set(qn("w:val"), "single")
    rpr.append(color); rpr.append(under); run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    run.append(t); link.append(run)
    par._p.append(link)


def add_runs(par, text):
    """Render inline **bold**, `code` and [links](url) into an existing paragraph."""
    for piece in re.split(r"(\[[^\]]+\]\([^)\s]+\)|\*\*.+?\*\*|`[^`]+`)", text):
        if not piece:
            continue
        m = LINK.fullmatch(piece)
        if m:
            add_link(par, m.group(1), m.group(2))
        elif piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(9.5)
            r.font.color.rgb = MONO_COLOR
        else:
            par.add_run(piece)


def is_table_row(line):
    return line.startswith("|") and line.endswith("|")


def split_row(line):
    return [c.strip() for c in line.strip("|").split("|")]


def is_divider(cells):
    return all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c)


def emit_table(doc, rows):
    """rows[0] is the header; a divider row has already been dropped."""
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            par = cells[j].paragraphs[0]
            add_runs(par, text)
            # Numbers read far better right-aligned; the first column is a label.
            if j and re.fullmatch(r"[−+-]?[\d.,]+", text):
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in par.runs:
                r.font.size = Pt(9.5)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def emit_code(doc, lines):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(18)
    par.paragraph_format.space_after = Pt(10)
    run = par.add_run("\n".join(lines))
    run.font.name = CODE_FONT
    run.font.size = Pt(8.5)


def convert(src, dst):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = open(src, encoding="utf-8").read().splitlines()
    i, para, table, code = 0, [], [], None

    def flush_para():
        if para:
            # The source is hard-wrapped; join it back into real paragraphs.
            add_runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    def flush_table():
        if table:
            emit_table(doc, list(table))
            table.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if code is not None:                      # inside a fenced block
            if stripped.startswith("```"):
                emit_code(doc, code)
                code = None
            else:
                code.append(line)
            i += 1
            continue

        if stripped.startswith("```"):
            flush_para(); flush_table()
            code = []
            i += 1
            continue

        if is_table_row(stripped):
            flush_para()
            cells = split_row(stripped)
            if not is_divider(cells):
                table.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            flush_para()
        elif stripped.startswith("#"):
            flush_para()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped.lstrip("# ").strip()
            # Headings take the text only: Word styles them already, and a
            # literal `backtick` in a heading just looks like a typo.
            title = re.sub(r"[`*]", "", title)
            doc.add_heading(title, level=min(level, 4))
        elif re.fullmatch(r"-{3,}", stripped):
            flush_para()                          # horizontal rule: just a break
        elif stripped.startswith(("- ", "* ")):
            flush_para()
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"\d+\.\s", stripped):
            flush_para()
            add_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\d+\.\s+", "", stripped))
        else:
            para.append(stripped)
        i += 1

    flush_para(); flush_table()
    if code:
        emit_code(doc, code)
    doc.save(dst)
    return dst


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "MEMO.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "MEMO.docx"
    print("wrote", convert(src, dst))
