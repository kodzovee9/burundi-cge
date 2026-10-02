"""Build the data-collection workbook for the Burundi 2019-2024 backcast.

    python3 tools/make_data_template.py

Writes pymodel/data/burundi-actuals-2019-2024-TEMPLATE.xlsx with five sheets:
README, Inputs (series that replace the model's 2020-24 projections), Targets
(outcomes the model's outputs are tested against), Sectors (value added by the
model's 19 sectors) and Trade (exports/imports by the model's 19 commodities).
Value cells are left empty for Kodzovi to fill; the 2019 "model" columns are
pre-filled from the SAM so the base year can be checked against actuals too.
"""
from __future__ import annotations

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "pymodel"))

from gemcore.database import load_database, bdi2019_data2          # noqa: E402
from gemcore.calibration import calibrate                          # noqa: E402

DATA = os.path.join(ROOT, "model", "user-files", "bdi2019", "bdi2019-data.xlsx")
OUT = os.path.join(ROOT, "pymodel", "data",
                   "burundi-actuals-2019-2024-TEMPLATE.xlsx")

YEARS = [str(y) for y in range(2019, 2026)]
HDR = PatternFill("solid", fgColor="1F4E78")
HDRF = Font(bold=True, color="FFFFFF")
FILL_A = PatternFill("solid", fgColor="F8CBAD")   # essential
FILL_B = PatternFill("solid", fgColor="FFE699")   # important
FILL_C = PatternFill("solid", fgColor="E2EFDA")   # nice to have
INPUT = PatternFill("solid", fgColor="FFFFFF")
MODEL = PatternFill("solid", fgColor="DDEBF7")    # pre-filled from model
WRAP = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------------------
# Series definitions.  (id, variable, description, model parameter, unit,
#                       priority, ideal source)
# ---------------------------------------------------------------------------
INPUTS = [
    ("IN01", "Official exchange rate", "Annual average, BIF per USD",
     "EXR (numeraire scaling)", "BIF/USD", "A",
     "BRB (Banque de la République du Burundi) statistics; IMF IFS"),
    ("IN02", "Parallel-market exchange rate", "Annual average, BIF per USD. "
     "The single most important series: the model's FX premium is currently "
     "held flat at 2.5 (150%) through 2027.",
     "prexrindex0(t) = IN02/IN01 ÷ 2019 value", "BIF/USD", "A",
     "IMF Article IV staff reports (cite the premium); BRB; press/market "
     "surveys; World Bank Burundi Economic Update"),
    ("IN03", "FX premium (derived)", "= IN02 / IN01. Fill only if you have the "
     "premium directly rather than the two rates.",
     "prexrindex0", "ratio", "A", "IMF Article IV"),
    ("IN04", "Real GDP growth", "Confirm the workbook's 2020-24 values "
     "(0.34, 3.12, 1.83, 2.66, 3.53%) are actuals; correct if not.",
     "gdpgrw(t) — PINS real GDP in the calibration pass", "%", "A",
     "ISTEEBU comptes nationaux; IMF WEO; World Bank WDI"),
    ("IN05", "Fuel import price", "Index or USD/bbl. Refined petroleum is a "
     "quota good carrying a model-generated 273% rent; the 2021-22 spike matters.",
     "pwmindex('c-refpet',t)", "index 2019=1 or USD/bbl", "A",
     "World Bank Pink Sheet (Brent); IMF Primary Commodity Prices; BRB import "
     "unit values"),
    ("IN06", "Import price index, all goods", "USD terms, 2019 = 1",
     "pwmindex(c,t)", "index", "B",
     "BRB / ISTEEBU import unit values; IMF WEO deflators"),
    ("IN07", "Export price index", "Coffee, tea, gold dominate; USD terms",
     "pweindex(c,t)", "index", "B",
     "ICO (coffee); World Bank Pink Sheet (tea, gold); BRB"),
    ("IN08", "Government consumption", "Held flat at 10.7% of GDP from 2020 "
     "in the workbook.", "govspndgdp0('congov',t)", "% of GDP", "A",
     "IMF Article IV fiscal tables; Ministry of Finance; BRB"),
    ("IN09", "Government capital spending", "Public investment. The workbook "
     "path has a discontinuity in 2020 (checklist C1).",
     "govspndgdp0('f-capgov',t)", "% of GDP", "A",
     "IMF Article IV fiscal tables; Ministry of Finance"),
    ("IN10", "Tax revenue", "Total tax revenue", "govrecgdp0(tax,t)",
     "% of GDP", "B", "IMF Article IV; OBR (Office Burundais des Recettes)"),
    ("IN11", "Grants to government", "External grants",
     "govrecgdp0('trgovrow',t)", "% of GDP", "B", "IMF Article IV"),
    ("IN12", "Government domestic financing", "Net domestic borrowing",
     "govrecgdp0('netdomfin',t)", "% of GDP", "B", "IMF Article IV; BRB"),
    ("IN13", "Government external financing", "Net external borrowing",
     "govrecgdp0('netforfingov',t)", "% of GDP", "B", "IMF Article IV"),
    ("IN14", "Population, total", "Mid-year", "pop0(h,t)", "thousands", "B",
     "ISTEEBU projections; UN World Population Prospects"),
    ("IN15", "Population, working age (15-64)", "Mid-year",
     "pop0('agelab',t)", "thousands", "B", "UN WPP; ILOSTAT"),
    ("IN16", "Import quota / licensing changes", "Any change to fuel or "
     "chemicals import quotas or FX allocation for imports. Text is fine.",
     "qmbarindex0(c,t)", "text / index", "C",
     "Ministry of Trade; BRB FX allocation circulars; IMF Article IV"),
    ("IN17", "Interest rate, government domestic debt", "Effective rate",
     "gintrat0(t)", "%", "C", "BRB; IMF DSA"),
    ("IN18", "Interest rate, external debt", "Effective rate",
     "fintrat0(ins2,t)", "%", "C", "IMF/WB Debt Sustainability Analysis"),
]

TARGETS = [
    ("TG01", "Nominal GDP", "Current prices", "GDPMP", "BIF billions", "A",
     "ISTEEBU comptes nationaux; IMF WEO"),
    ("TG02", "CPI inflation", "Annual average", "CPI growth", "%", "A",
     "ISTEEBU IPC; IMF IFS"),
    ("TG03", "Exports of goods", "BoP basis", "Σ PWE·QE", "USD millions", "A",
     "BRB balance of payments; IMF BoP statistics"),
    ("TG04", "Exports of goods and services", "BoP basis", "", "USD millions",
     "B", "BRB balance of payments"),
    ("TG05", "Imports of goods", "BoP basis", "Σ PWM·QM", "USD millions", "A",
     "BRB balance of payments"),
    ("TG06", "Imports of goods and services", "BoP basis", "", "USD millions",
     "B", "BRB balance of payments"),
    ("TG07", "Current account balance", "Including grants",
     "SAVF (sign reversed)", "% of GDP", "A", "BRB; IMF Article IV"),
    ("TG08", "External public debt", "Government and guaranteed",
     "FDEBT('govz')", "USD millions or % GDP", "A",
     "IMF Article IV / DSA; World Bank IDS"),
    ("TG09", "External debt, total", "Public + private",
     "FDEBT total", "USD millions or % GDP", "B", "World Bank IDS; IMF DSA"),
    ("TG10", "Government domestic debt", "", "GDEBT", "% of GDP", "B",
     "IMF Article IV; BRB"),
    ("TG11", "Gross fixed capital formation", "Total", "PrvFixInv+GovFixInv",
     "% of GDP", "A", "ISTEEBU; World Bank WDI"),
    ("TG12", "Private consumption", "Households + NPISH", "PrvCon",
     "% of GDP", "A", "ISTEEBU; World Bank WDI"),
    ("TG13", "Government consumption", "Same as IN08 — a target when the "
     "closure makes it endogenous", "GovCon", "% of GDP", "B", "as IN08"),
    ("TG14", "Unemployment rate", "Total, if any estimate exists",
     "UERAT", "%", "C", "ILOSTAT modelled estimates; ISTEEBU survey"),
    ("TG15", "Wages", "Nominal wage index or public-sector wage bill growth",
     "WF", "index or %", "C", "ISTEEBU; Ministry of Finance"),
    ("TG16", "FDI inflows", "Net", "INVVALF", "USD millions", "B",
     "BRB BoP; UNCTAD"),
    ("TG17", "Remittances", "Inflows", "TRII(row,h)", "USD millions", "B",
     "BRB BoP; World Bank remittance data"),
    ("TG18", "Official reserves", "End-year; or months of imports", "drf",
     "USD millions", "B", "BRB; IMF Article IV"),
    ("TG19", "Coffee exports", "Value and volume", "QE('c-agr')/('c-food')",
     "USD mn; tonnes", "B", "BRB; ICO; OCIBU/ARFIC"),
    ("TG20", "Mining exports (gold, rare earths, nickel)", "Value. Relevant to "
     "the combi scenario later.", "QE('c-min')", "USD millions", "B",
     "BRB; EITI Burundi; ISTEEBU"),
]

SECTOR_LABELS = {
    "a-agr": "Agriculture, livestock, forestry, fishing",
    "a-min": "Mining and quarrying",
    "a-food": "Food, beverages, tobacco",
    "a-texwapp": "Textiles and apparel",
    "a-woodpaper": "Wood and paper products",
    "a-chemplast": "Chemicals and plastics (incl. refined petroleum?)",
    "a-prodminnmet": "Non-metallic mineral products (cement etc.)",
    "a-met": "Metals and metal products",
    "a-oman": "Other manufacturing",
    "a-elect": "Electricity",
    "a-water": "Water",
    "a-construc": "Construction",
    "a-trade": "Trade (wholesale, retail)",
    "a-transp": "Transport and communications",
    "a-hotelrest": "Hotels and restaurants",
    "a-admpub": "Public administration",
    "a-edu": "Education",
    "a-health": "Health",
    "a-oser": "Other services",
}


def header(ws, cols, widths):
    ws.append(cols)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j)
        cell.fill, cell.font, cell.alignment = HDR, HDRF, WRAP
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32


def prio_fill(p):
    return {"A": FILL_A, "B": FILL_B, "C": FILL_C}[p]


def series_sheet(ws, rows):
    cols = (["ID", "Variable", "Description", "Model parameter / variable",
             "Unit", "Priority"] + YEARS
            + ["Ideal source", "Source used", "Notes"])
    widths = [7, 30, 46, 30, 16, 9] + [10] * len(YEARS) + [44, 30, 30]
    header(ws, cols, widths)
    for r in rows:
        ws.append(list(r[:6]) + [None] * len(YEARS) + [r[6], None, None])
        i = ws.max_row
        ws.cell(row=i, column=6).fill = prio_fill(r[5])
        for j in range(7, 7 + len(YEARS)):
            ws.cell(row=i, column=j).fill = INPUT
        for j in (2, 3, 4, 7 + len(YEARS)):
            ws.cell(row=i, column=j).alignment = WRAP


def build():
    cal = calibrate(load_database(DATA, data2_hook=bdi2019_data2))
    S, sam = cal.db.sets, cal.db.pars["sam"]
    A, C, F = S["a"], S["c"], S["f"]
    ROW = S["insrow"]
    va = {a: sum(sam.get((f, a), 0.0) for f in F) for a in A}
    va_tot = sum(va.values())
    exp_c = {c: sum(sam.get((c, r), 0.0) for r in ROW) for c in C}
    imp_c = {c: sum(sam.get((r, c), 0.0) for r in ROW) for c in C}
    gdp = cal.GDPMP00

    wb = Workbook()

    # ---- README -----------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Burundi CGE — data collection for the 2019-2024 backcast", True),
        ("", False),
        ("PURPOSE", True),
        ("The Python CGE model is being validated against Burundi's actual "
         "2019-2024 outcomes rather than against the GAMS projection. The model "
         "keeps the 2019 SAM as base. For 2020-2024 it needs (a) actual values "
         "of the series it takes as exogenous INPUTS, and (b) actual OUTCOMES "
         "to test its outputs against.", False),
        ("", False),
        ("SHEETS", True),
        ("Inputs   — series that replace the model's 2020-24 projections. "
         "Filling these changes what the model is fed.", False),
        ("Targets  — outcomes the model's outputs are compared to. Filling these "
         "changes nothing in the model; they are the test.", False),
        ("Sectors  — value added by the model's 19 sectors. The 2019 'model "
         "share' column is pre-filled from the SAM for you to check.", False),
        ("Trade    — exports and imports by the model's 19 commodities; 2019 "
         "model values pre-filled.", False),
        ("", False),
        ("PERIOD", True),
        ("2019 is the base year (already in the SAM; use it to check the SAM). "
         "2020-2024 are the actuals we need. 2025 if an estimate exists.", False),
        ("", False),
        ("PRIORITY", True),
        ("A (red)    essential — the backcast cannot be judged without it", False),
        ("B (yellow) important — improves the test materially", False),
        ("C (green)  nice to have", False),
        ("Fill the A rows first, across all five sheets. IN02 (parallel exchange "
         "rate), IN04 (GDP growth confirmation), IN05 (fuel price), TG02 (CPI), "
         "TG03/TG05 (trade) and the Sectors sheet are the ones that decide "
         "whether this works.", False),
        ("", False),
        ("CONVENTIONS", True),
        ("Percentages as percent, not decimals (3.5 for 3.5%).", False),
        ("Currency: state the unit in the Unit column if it differs from the "
         "one given; BIF billions and USD millions preferred.", False),
        ("Indices: 2019 = 1.00 unless stated.", False),
        ("Leave a cell EMPTY if the value is not available — do not enter 0.", False),
        ("Put the source actually used in 'Source used' (report name and table, "
         "or URL). If a value is an estimate or projection rather than an "
         "actual, say so in Notes.", False),
        ("Annual averages unless the description says end-year.", False),
        ("", False),
        ("WHAT HAPPENS NEXT", True),
        ("Return the workbook as is. A loader is written around this exact "
         "layout, so do not rename sheets, move columns, or delete ID rows; "
         "adding rows at the bottom of a sheet is fine.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).alignment = WRAP
        if bold:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    # ---- Inputs / Targets -------------------------------------------------
    series_sheet(wb.create_sheet("Inputs"), INPUTS)
    series_sheet(wb.create_sheet("Targets"), TARGETS)

    # ---- Sectors ----------------------------------------------------------
    ws = wb.create_sheet("Sectors")
    cols = (["Model sector", "Description", "National-accounts category "
             "you mapped it to", "2019 model VA share (%)", "Unit of your data"]
            + [f"{y} actual" for y in YEARS] + ["Source used", "Notes"])
    header(ws, cols, [16, 40, 34, 14, 22] + [11] * len(YEARS) + [30, 30])
    for a in A:
        ws.append([a, SECTOR_LABELS.get(a, ""), None,
                   round(100 * va[a] / va_tot, 3),
                   "VA, BIF bn current  OR  real growth %  OR  share %"]
                  + [None] * len(YEARS) + [None, None])
        i = ws.max_row
        ws.cell(row=i, column=4).fill = MODEL
        for j in (2, 3, 5):
            ws.cell(row=i, column=j).alignment = WRAP
    ws.append(["TOTAL", "Sum of the above", None, 100.0, "GDP at factor cost"]
              + [None] * len(YEARS) + [None, None])
    ws.cell(row=ws.max_row, column=4).fill = MODEL
    ws.append([])
    ws.append(["Fill whatever form your national accounts give — nominal value "
               "added, real growth, or shares. State which in the Unit column. "
               "Coarser categories (e.g. 'manufacturing' covering a-food through "
               "a-oman) are fine: write the same category name against each model "
               "sector it covers and give the total once."])
    ws.cell(row=ws.max_row, column=1).alignment = WRAP
    ws.merge_cells(start_row=ws.max_row, start_column=1,
                   end_row=ws.max_row, end_column=8)
    ws.row_dimensions[ws.max_row].height = 48

    # ---- Trade ------------------------------------------------------------
    ws = wb.create_sheet("Trade")
    cols = (["Model commodity", "Description", "Flow",
             "2019 model value (SAM units, 10 bn BIF)", "2019 model % of GDP",
             "Unit of your data"]
            + [f"{y} actual" for y in YEARS] + ["Source used", "Notes"])
    header(ws, cols, [16, 38, 9, 16, 12, 22] + [11] * len(YEARS) + [30, 30])
    for c in C:
        for flow, val in (("exports", exp_c[c]), ("imports", imp_c[c])):
            if val <= 0:
                continue
            ws.append([c, SECTOR_LABELS.get(c.replace("c-", "a-"), ""), flow,
                       round(val, 4), round(100 * val / gdp, 3),
                       "USD mn  OR  BIF bn"]
                      + [None] * len(YEARS) + [None, None])
            i = ws.max_row
            ws.cell(row=i, column=4).fill = MODEL
            ws.cell(row=i, column=5).fill = MODEL
            ws.cell(row=i, column=2).alignment = WRAP
    ws.append([])
    ws.append(["Priority C overall, except c-refpet and c-chemplast imports "
               "(the quota goods) and c-min / c-agr exports, which are B. "
               "HS-chapter or SITC data from BRB or UN Comtrade can be mapped to "
               "these 19 commodities; note the mapping used."])
    ws.cell(row=ws.max_row, column=1).alignment = WRAP
    ws.merge_cells(start_row=ws.max_row, start_column=1,
                   end_row=ws.max_row, end_column=8)
    ws.row_dimensions[ws.max_row].height = 40

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print("wrote", build())
